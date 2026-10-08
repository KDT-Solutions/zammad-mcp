"""
Textextraktion fuer Ticket-Anhaenge (PDF, XLSX, DOCX, CSV, TXT).

Dieses Modul laeuft bewusst in einem eigenen Python-Prozess (siehe run_extraction
im Paket): der Hauptprozess schreibt die Datei-Bytes auf stdin, das Ergebnis kommt
als JSON auf stdout zurueck. So kann eine haengende oder speicherhungrige Datei
(kaputte PDF, Zip-Bombe, boesartiger Font) hart per Timeout beendet werden, ohne den
MCP-Server mitzureissen. Deshalb hier nur absolute Imports von Stdlib und Drittpaketen,
nichts aus dem zammad_mcp-Paket.

PDF-Ablauf pro Seite:
1. pdfminer.six (Layout-Analyse, Zeilen nach y gruppiert, innerhalb nach x sortiert)
2. Ist der Text Zeichensalat ((cid:..), Steuerzeichen, U+FFFD ...), werden die
   Content-Streams selbst interpretiert und die Glyph-IDs ueber die cmap-Tabelle der
   eingebetteten Fonts (FontFile2, via fontTools) zurueck nach Unicode uebersetzt.
3. Liefert das nichts Brauchbares: GID + 29 = ASCII (Standard-Glyph-Reihenfolge
   vieler TrueType-Subsets, z.B. von WPS Office erzeugte PDFs).
"""

import csv
import io
import json
import math
import re
import sys
import unicodedata
import zipfile

SUPPORTED_KINDS = ("pdf", "xlsx", "docx", "csv", "txt")

EXTENSION_KINDS = {".pdf": "pdf", ".xlsx": "xlsx", ".xlsm": "xlsx", ".docx": "docx", ".csv": "csv", ".txt": "txt"}
CONTENT_TYPE_KINDS = {
    "application/pdf": "pdf",
    "application/x-pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel.sheet.macroenabled.12": "xlsx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/csv": "csv",
    "application/csv": "csv",
    "text/plain": "txt",
}

# Obergrenze fuer den extrahierten Text (Schutz vor Speicher-Explosion bei
# riesigen Tabellen); bei Ueberschreitung wird abgeschnitten und das gemeldet.
MAX_TEXT_CHARS = 5_000_000
# Maximale entpackte Groesse von XLSX/DOCX (Zip-Bomben-Schutz)
MAX_ZIP_UNCOMPRESSED = 300 * 1024 * 1024

GARBLED_THRESHOLD = 0.2
_CID_RE = re.compile(r"\(cid:\d+\)")


class ExtractionLimit(Exception):
    """Ein Sicherheitslimit wurde ueberschritten (limit = 'pages', 'size', ...)."""

    def __init__(self, limit: str, message: str):
        super().__init__(message)
        self.limit = limit


def detect_kind(filename: str = "", content_type: str = "", data: bytes | None = None) -> str | None:
    """Dateityp aus Endung, Content-Type und (falls vorhanden) Magic Bytes bestimmen."""
    name = (filename or "").lower()
    kind = next((k for ext, k in EXTENSION_KINDS.items() if name.endswith(ext)), None)
    if kind is None:
        kind = CONTENT_TYPE_KINDS.get((content_type or "").split(";")[0].strip().lower())
    if data is not None:
        if data[:5] == b"%PDF-" or (kind == "pdf" and b"%PDF-" in data[:1024]):
            return "pdf"
        if kind == "pdf":
            return None
        if kind in ("xlsx", "docx") and not data.startswith(b"PK"):
            return None
    return kind


# ---------------------------------------------------------------------------
# Hilfsfunktionen
# ---------------------------------------------------------------------------

def garbled_ratio(text: str) -> float:
    """Anteil unbrauchbarer Zeichen: (cid:N)-Tokens, Steuerzeichen, Private Use, U+FFFD."""
    cid_chars = sum(len(m) for m in _CID_RE.findall(text))
    rest = _CID_RE.sub("", text)
    bad = 0
    total = cid_chars
    for ch in rest:
        if ch.isspace():
            continue
        total += 1
        if ch == "�" or unicodedata.category(ch) in ("Cc", "Co", "Cn", "Cs"):
            bad += 1
    if total == 0:
        return 0.0
    return (cid_chars + bad) / total


def is_garbled(text: str) -> bool:
    return garbled_ratio(text) > GARBLED_THRESHOLD


def _join_line(fragments: list[tuple[float, float, str, float]]) -> str:
    """Fragmente (x0, x1, text, fontsize) einer Zeile nach x sortiert zusammensetzen.
    Grosse Luecken (Tabellenspalten) werden als drei Leerzeichen ausgegeben."""
    out = ""
    prev_x1 = None
    for x0, x1, text, size in sorted(fragments, key=lambda f: f[0]):
        if not text:
            continue
        if prev_x1 is not None and out and not out.endswith(" ") and not text.startswith(" "):
            gap = x0 - prev_x1
            size = max(size, 1.0)
            if gap > 2.0 * size:
                out += "   "
            elif gap > 0.15 * size:
                out += " "
        out += text
        prev_x1 = x1 if prev_x1 is None else max(prev_x1, x1)
    return out.rstrip()


def _group_lines(fragments: list[tuple[float, float, float, str, float]]) -> str:
    """Fragmente (y, x0, x1, text, fontsize) zu Zeilen gruppieren: gleiche y-Koordinate
    (Toleranz halbe Schriftgroesse) = eine Zeile, Zeilen von oben nach unten."""
    lines: list[tuple[float, float, list]] = []  # (y, tol, fragments)
    for y, x0, x1, text, size in sorted(fragments, key=lambda f: (-f[0], f[1])):
        tol = max(size, 1.0) * 0.5
        if lines and abs(lines[-1][0] - y) <= max(tol, lines[-1][1]):
            lines[-1][2].append((x0, x1, text, size))
        else:
            lines.append((y, tol, [(x0, x1, text, size)]))
    return "\n".join(_join_line(frags) for _, _, frags in lines).strip("\n")


# ---------------------------------------------------------------------------
# PDF: Primaer pdfminer.six
# ---------------------------------------------------------------------------

def _pdfminer_pages(data: bytes, max_pages: int) -> list[str]:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LAParams, LTTextLine, LTLayoutContainer

    def collect(obj, out):
        if isinstance(obj, LTTextLine):
            text = obj.get_text().replace("\n", " ").rstrip()
            if text.strip():
                out.append((obj.y0, obj.x0, obj.x1, text, obj.height))
            return
        if isinstance(obj, LTLayoutContainer) or hasattr(obj, "__iter__"):
            try:
                for child in obj:
                    collect(child, out)
            except TypeError:
                pass

    pages = []
    for layout in extract_pages(io.BytesIO(data), laparams=LAParams(all_texts=True), maxpages=max_pages):
        frags: list = []
        collect(layout, frags)
        pages.append(_group_lines(frags))
    return pages


# ---------------------------------------------------------------------------
# PDF: Fallback - Content-Streams selbst interpretieren, Glyph-IDs ueber Fonts mappen
# ---------------------------------------------------------------------------

def _mult(m1, m2):
    a1, b1, c1, d1, e1, f1 = m1
    a2, b2, c2, d2, e2, f2 = m2
    return (
        a1 * a2 + b1 * c2, a1 * b2 + b1 * d2,
        c1 * a2 + d1 * c2, c1 * b2 + d1 * d2,
        e1 * a2 + f1 * c2 + e2, e1 * b2 + f1 * d2 + f2,
    )


_IDENTITY = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)


def _num(x, default=0.0) -> float:
    try:
        return float(x)
    except Exception:
        return default


def _parse_tounicode(data: bytes) -> dict[int, str]:
    """Minimaler Parser fuer ToUnicode-CMaps (bfchar / bfrange)."""
    text = data.decode("latin-1", "replace")
    result: dict[int, str] = {}

    def hex_to_str(h: str) -> str:
        h = re.sub(r"\s", "", h)
        if len(h) % 4:
            h = h.rjust(len(h) + 4 - len(h) % 4, "0")
        try:
            return bytes.fromhex(h).decode("utf-16-be", "replace")
        except ValueError:
            return ""

    for block in re.findall(r"beginbfchar(.*?)endbfchar", text, re.S):
        for src, dst in re.findall(r"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f\s]*)>", block):
            result[int(src, 16)] = hex_to_str(dst)
    for block in re.findall(r"beginbfrange(.*?)endbfrange", text, re.S):
        for lo, hi, rest in re.findall(r"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*(\[[^\]]*\]|<[0-9A-Fa-f\s]*>)", block):
            lo_i, hi_i = int(lo, 16), int(hi, 16)
            if hi_i - lo_i > 65535:
                continue
            if rest.startswith("["):
                for i, h in enumerate(re.findall(r"<([0-9A-Fa-f\s]*)>", rest)):
                    if lo_i + i > hi_i:
                        break
                    result[lo_i + i] = hex_to_str(h)
            else:
                base = hex_to_str(rest.strip("<>"))
                if not base:
                    continue
                for i in range(hi_i - lo_i + 1):
                    result[lo_i + i] = base[:-1] + chr(min(ord(base[-1]) + i, 0x10FFFF))
    return result


def _gid_to_unicode_from_font(font_bytes: bytes) -> dict[int, str]:
    """Reverse-Map GID -> Unicode aus der cmap-Tabelle (und notfalls den Glyph-Namen
    der post-Tabelle) eines eingebetteten TrueType/OpenType-Fonts."""
    from fontTools.ttLib import TTFont
    from fontTools import agl

    font = TTFont(io.BytesIO(font_bytes), lazy=True, fontNumber=0)
    rev: dict[int, str] = {}
    try:
        glyph_order = font.getGlyphOrder()
    except Exception:
        glyph_order = []
    name_to_gid = {name: i for i, name in enumerate(glyph_order)}

    def put(gid, ch, weak=False):
        if gid is None:
            return
        cur = rev.get(gid)
        if cur is None or (not weak and unicodedata.category(cur) == "Co" and unicodedata.category(ch) != "Co"):
            rev[gid] = ch

    try:
        cmap = font["cmap"]
        tables = list(cmap.tables)
    except Exception:
        tables = []
    # Unicode-Tabellen zuerst, Symbol-Tabelle (3,0) mit F0xx-Codes nur als Notnagel
    for t in sorted(tables, key=lambda t: 0 if t.isUnicode() else 1):
        try:
            items = t.cmap.items()
        except Exception:
            continue
        for cp, name in items:
            gid = name_to_gid.get(name)
            if gid is None:
                try:
                    gid = font.getGlyphID(name)
                except Exception:
                    continue
            if t.isUnicode():
                if 0 <= cp <= 0x10FFFF:
                    put(gid, chr(cp))
            elif t.platformID == 3 and t.platEncID == 0 and 0xF020 <= cp <= 0xF0FF:
                put(gid, chr(cp - 0xF000), weak=True)

    # Glyph-Namen (post-Tabelle), z.B. 'A', 'uni00C4', 'adieresis'
    for gid, name in enumerate(glyph_order):
        if gid in rev or not name or name.startswith((".", "glyph")) or name in (".notdef",):
            continue
        try:
            ch = agl.toUnicode(name)
        except Exception:
            ch = ""
        if ch:
            rev[gid] = ch
    return rev


def _font_file_bytes(descriptor) -> bytes | None:
    if descriptor is None:
        return None
    for key in ("/FontFile2", "/FontFile3"):
        ff = descriptor.get(key)
        if ff is not None:
            try:
                return ff.get_object().get_data()
            except Exception:
                return None
    return None


class _FontDecoder:
    """Dekodiert Strings eines PDF-Fonts zu Unicode und berechnet Glyph-Breiten."""

    def __init__(self, font_obj):
        f = font_obj.get_object()
        self.subtype = str(f.get("/Subtype", ""))
        self.two_byte = self.subtype == "/Type0"
        self.cid_to_gid: bytes | None = None  # None = Identity
        self.widths: dict[int, float] = {}
        self.default_width = 1000.0 if self.two_byte else 500.0
        self.tounicode: dict[int, str] = {}
        self.gid2uni: dict[int, str] = {}
        self.code_to_gid: dict[int, int] = {}
        self.has_font_file = False

        tu = f.get("/ToUnicode")
        if tu is not None:
            try:
                self.tounicode = _parse_tounicode(tu.get_object().get_data())
            except Exception:
                self.tounicode = {}

        descriptor = None
        if self.two_byte:
            desc_fonts = f.get("/DescendantFonts")
            cid_font = None
            try:
                cid_font = desc_fonts.get_object()[0].get_object()
            except Exception:
                cid_font = None
            if cid_font is not None:
                descriptor = cid_font.get("/FontDescriptor")
                descriptor = descriptor.get_object() if descriptor is not None else None
                self.default_width = _num(cid_font.get("/DW", 1000), 1000.0)
                self._parse_w(cid_font.get("/W"))
                c2g = cid_font.get("/CIDToGIDMap")
                if c2g is not None:
                    c2g = c2g.get_object()
                    if hasattr(c2g, "get_data"):
                        try:
                            self.cid_to_gid = c2g.get_data()
                        except Exception:
                            self.cid_to_gid = None
        else:
            descriptor = f.get("/FontDescriptor")
            descriptor = descriptor.get_object() if descriptor is not None else None
            first = int(_num(f.get("/FirstChar", 0)))
            widths = f.get("/Widths")
            if widths is not None:
                try:
                    for i, w in enumerate(widths.get_object()):
                        self.widths[first + i] = _num(w, self.default_width)
                except Exception:
                    pass

        font_bytes = _font_file_bytes(descriptor)
        if font_bytes:
            try:
                self.gid2uni = _gid_to_unicode_from_font(font_bytes)
                self.has_font_file = True
                if not self.two_byte:
                    self._simple_code_to_gid(font_bytes)
            except Exception:
                self.gid2uni = {}

    def _parse_w(self, w):
        if w is None:
            return
        try:
            arr = [x.get_object() if hasattr(x, "get_object") else x for x in w.get_object()]
        except Exception:
            return
        i = 0
        while i < len(arr):
            try:
                start = int(arr[i])
                nxt = arr[i + 1]
                if isinstance(nxt, list) or hasattr(nxt, "__iter__"):
                    for j, wv in enumerate(nxt):
                        self.widths[start + j] = _num(wv, self.default_width)
                    i += 2
                else:
                    end = int(nxt)
                    wv = _num(arr[i + 2], self.default_width)
                    if end - start <= 65535:
                        for c in range(start, end + 1):
                            self.widths[c] = wv
                    i += 3
            except Exception:
                break

    def _simple_code_to_gid(self, font_bytes: bytes):
        """Einfache TrueType-Fonts: Byte-Code -> GID ueber die (3,0)/(1,0)-cmap des Fonts."""
        from fontTools.ttLib import TTFont
        font = TTFont(io.BytesIO(font_bytes), lazy=True)
        try:
            order = {n: i for i, n in enumerate(font.getGlyphOrder())}
            for t in font["cmap"].tables:
                for cp, name in t.cmap.items():
                    gid = order.get(name)
                    if gid is None:
                        continue
                    if t.platformID == 3 and t.platEncID == 0 and 0xF000 <= cp <= 0xF0FF:
                        self.code_to_gid.setdefault(cp - 0xF000, gid)
                    elif t.platformID == 1 and t.platEncID == 0 and cp <= 0xFF:
                        self.code_to_gid.setdefault(cp, gid)
        except Exception:
            pass

    def codes(self, raw: bytes) -> list[int]:
        if self.two_byte:
            if len(raw) % 2:
                raw = raw + b"\x00"
            return [(raw[i] << 8) | raw[i + 1] for i in range(0, len(raw), 2)]
        return list(raw)

    def gid(self, code: int) -> int | None:
        if self.two_byte:
            if self.cid_to_gid is None:
                return code
            idx = code * 2
            if idx + 1 < len(self.cid_to_gid):
                return (self.cid_to_gid[idx] << 8) | self.cid_to_gid[idx + 1]
            return None
        return self.code_to_gid.get(code)

    def char(self, code: int, mode: str) -> str:
        gid = self.gid(code)
        # Embedded-Font zuerst: die ToUnicode-Map ist bei den Problem-PDFs genau das,
        # was fehlt oder kaputt ist.
        if gid is not None and gid in self.gid2uni:
            return self.gid2uni[gid]
        tu = self.tounicode.get(code)
        if tu and not is_garbled(tu):
            return tu
        if not self.two_byte:
            try:
                return bytes([code]).decode("cp1252")
            except UnicodeDecodeError:
                return chr(code) if code >= 0x20 else "�"
        if mode == "gid" and gid is not None:
            # Standard-Glyph-Reihenfolge: GID 3 = Leerzeichen, GID 36 = 'A' ...
            cp = gid + 29
            if 32 <= cp < 127:
                return chr(cp)
        return "�"

    def width(self, code: int) -> float:
        return self.widths.get(code, self.default_width)


def _resolve(obj):
    return obj.get_object() if hasattr(obj, "get_object") else obj


def _raw_bytes(s) -> bytes:
    if isinstance(s, bytes):
        return bytes(s)
    orig = getattr(s, "original_bytes", None)
    if orig is not None:
        return bytes(orig)
    return str(s).encode("latin-1", "replace")


def _decode_page_with_fonts(page, reader, mode: str, font_cache: dict) -> str:
    """Content-Stream einer Seite interpretieren und Text mit Positionen sammeln.
    mode='cmap': nur echte Unicode-Infos (Font-cmap, ToUnicode, Encoding)
    mode='gid':  zusaetzlich unbekannte GIDs ueber GID+29 raten."""
    from pypdf.generic import ContentStream

    fragments: list = []

    def get_font(resources, name):
        try:
            fonts = _resolve(_resolve(resources).get("/Font"))
            ref = fonts.get(name)
        except Exception:
            return None
        if ref is None:
            return None
        key = getattr(ref, "idnum", None) or id(_resolve(ref))
        if key not in font_cache:
            try:
                font_cache[key] = _FontDecoder(ref)
            except Exception:
                font_cache[key] = None
        return font_cache[key]

    def run(stream_ops, resources, ctm, depth):
        gs_stack = []
        tm = tlm = _IDENTITY
        font = None
        size = 0.0
        tc = tw = ts = 0.0
        th = 1.0
        tl = 0.0

        def show(items):
            nonlocal tm
            if font is None:
                return
            text = ""
            trm = _mult((size * th, 0, 0, size, 0, ts), _mult(tm, ctm))
            x_start, y = trm[4], trm[5]
            eff_size = math.hypot(trm[2], trm[3]) or 1.0
            for item in items:
                if isinstance(item, (int, float)):
                    adj = _num(item)
                    tx = -adj / 1000.0 * size * th
                    tm = _mult((1, 0, 0, 1, tx, 0), tm)
                    if adj < -250 and text and not text.endswith(" "):
                        text += " "
                    continue
                raw = _raw_bytes(item)
                for code in font.codes(raw):
                    text += font.char(code, mode)
                    tx = (font.width(code) / 1000.0 * size + tc + (tw if (not font.two_byte and code == 32) else 0.0)) * th
                    tm = _mult((1, 0, 0, 1, tx, 0), tm)
            x_end = _mult((size * th, 0, 0, size, 0, ts), _mult(tm, ctm))[4]
            if text.strip():
                fragments.append((y, min(x_start, x_end), max(x_start, x_end), text, eff_size))

        for operands, op in stream_ops:
            op = op.decode("latin-1") if isinstance(op, bytes) else str(op)
            try:
                if op == "q":
                    gs_stack.append(ctm)
                elif op == "Q":
                    if gs_stack:
                        ctm = gs_stack.pop()
                elif op == "cm" and len(operands) == 6:
                    ctm = _mult(tuple(_num(x) for x in operands), ctm)
                elif op == "BT":
                    tm = tlm = _IDENTITY
                elif op == "Tf" and len(operands) == 2:
                    font = get_font(resources, operands[0])
                    size = _num(operands[1])
                elif op == "Tc":
                    tc = _num(operands[0])
                elif op == "Tw":
                    tw = _num(operands[0])
                elif op == "Tz":
                    th = _num(operands[0], 100.0) / 100.0
                elif op == "TL":
                    tl = _num(operands[0])
                elif op == "Ts":
                    ts = _num(operands[0])
                elif op in ("Td", "TD") and len(operands) == 2:
                    tx, ty = _num(operands[0]), _num(operands[1])
                    if op == "TD":
                        tl = -ty
                    tlm = _mult((1, 0, 0, 1, tx, ty), tlm)
                    tm = tlm
                elif op == "Tm" and len(operands) == 6:
                    tlm = tm = tuple(_num(x) for x in operands)
                elif op == "T*":
                    tlm = _mult((1, 0, 0, 1, 0, -tl), tlm)
                    tm = tlm
                elif op == "Tj" and operands:
                    show([operands[0]])
                elif op == "TJ" and operands:
                    show(list(operands[0]))
                elif op == "'" and operands:
                    tlm = _mult((1, 0, 0, 1, 0, -tl), tlm)
                    tm = tlm
                    show([operands[0]])
                elif op == '"' and len(operands) == 3:
                    tw, tc = _num(operands[0]), _num(operands[1])
                    tlm = _mult((1, 0, 0, 1, 0, -tl), tlm)
                    tm = tlm
                    show([operands[2]])
                elif op == "Do" and operands and depth < 5:
                    xobjs = _resolve(_resolve(resources).get("/XObject")) if resources is not None else None
                    xobj = _resolve(xobjs.get(operands[0])) if xobjs is not None else None
                    if xobj is not None and str(xobj.get("/Subtype")) == "/Form":
                        matrix = tuple(_num(x) for x in xobj.get("/Matrix", [1, 0, 0, 1, 0, 0]))
                        sub_res = xobj.get("/Resources", resources)
                        sub_ops = ContentStream(xobj, reader).operations
                        run(sub_ops, sub_res, _mult(matrix, ctm), depth + 1)
            except Exception:
                continue

    contents = page.get_contents()
    if contents is None:
        return ""
    run(contents.operations, page.get("/Resources"), _IDENTITY, 0)
    return _group_lines(fragments)


def extract_pdf(data: bytes, max_pages: int) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise ValueError("PDF ist passwortgeschuetzt")
        except Exception:
            raise ValueError("PDF ist passwortgeschuetzt")
    page_count = len(reader.pages)
    if page_count > max_pages:
        raise ExtractionLimit("pages", f"PDF hat {page_count} Seiten, erlaubt sind maximal {max_pages}")

    try:
        miner_pages = _pdfminer_pages(data, max_pages)
        miner_error = None
    except Exception as e:
        miner_pages = []
        miner_error = f"{type(e).__name__}: {e}"

    font_cache: dict = {}
    pages: list[str] = []
    methods: list[str] = []
    for i in range(page_count):
        text = miner_pages[i] if i < len(miner_pages) else ""
        method = "pdfminer"
        if miner_error is not None or is_garbled(text):
            page = reader.pages[i]
            try:
                cmap_text = _decode_page_with_fonts(page, reader, "cmap", font_cache)
            except Exception:
                cmap_text = ""
            if cmap_text.strip() and not is_garbled(cmap_text):
                text, method = cmap_text, "font_cmap"
            else:
                try:
                    gid_text = _decode_page_with_fonts(page, reader, "gid", font_cache)
                except Exception:
                    gid_text = ""
                if gid_text.strip() and garbled_ratio(gid_text) < min(garbled_ratio(text) if text.strip() else 1.0, GARBLED_THRESHOLD):
                    text, method = gid_text, "gid_offset_heuristic"
                elif miner_error is not None and cmap_text.strip():
                    text, method = cmap_text, "font_cmap"
        pages.append(text)
        methods.append(method)

    # Gesamtmethode = die "unsicherste" Methode, die auf irgendeiner Seite noetig war
    overall = "pdfminer"
    for m in ("font_cmap", "gid_offset_heuristic"):
        if m in methods:
            overall = m
    by_method: dict[str, list[int]] = {}
    for i, m in enumerate(methods, start=1):
        by_method.setdefault(m, []).append(i)

    sections = []
    parts = []
    pos = 0
    for i, t in enumerate(pages, start=1):
        header = f"--- Seite {i} ---\n"
        chunk = header + t + "\n\n"
        sections.append({"offset": pos, "page": i})
        parts.append(chunk)
        pos += len(chunk)
    result = {
        "text": "".join(parts).rstrip("\n") + "\n",
        "pages": page_count,
        "extraction_method": overall,
        "pages_by_method": by_method,
        "sections": sections,
    }
    if miner_error:
        result["pdfminer_error"] = miner_error
    return result


# ---------------------------------------------------------------------------
# XLSX / DOCX / CSV / TXT
# ---------------------------------------------------------------------------

def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            total = sum(i.file_size for i in zf.infolist())
    except zipfile.BadZipFile:
        raise ValueError("Datei ist kein gueltiges Office-Dokument (ZIP-Container defekt)")
    if total > MAX_ZIP_UNCOMPRESSED:
        raise ExtractionLimit("size", f"Entpackte Groesse {total} Bytes ueber dem Limit von {MAX_ZIP_UNCOMPRESSED} Bytes")


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def extract_xlsx(data: bytes) -> dict:
    _check_zip(data)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, sections, pos, total = [], [], 0, 0
    try:
        for ws in wb.worksheets:
            header = f"--- Blatt: {ws.title} ---\n"
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [_cell_str(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    line = "\t".join(cells)
                    rows.append(line)
                    total += len(line) + 1
                    if total > MAX_TEXT_CHARS:
                        break
            chunk = header + "\n".join(rows) + "\n\n"
            sections.append({"offset": pos, "sheet": ws.title})
            parts.append(chunk)
            pos += len(chunk)
            if total > MAX_TEXT_CHARS:
                break
    finally:
        wb.close()
    return {"text": "".join(parts).rstrip("\n") + "\n", "sections": sections, "sheets": len(sections)}


def extract_docx(data: bytes) -> dict:
    _check_zip(data)
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))

    def table_lines(tbl) -> list[str]:
        out = []
        for row in tbl.rows:
            cells = []
            for cell in row.cells:
                t = cell.text.replace("\n", " ").strip()
                # verbundene Zellen liefert python-docx mehrfach
                if not cells or cells[-1] != t:
                    cells.append(t)
            if any(cells):
                out.append("\t".join(cells))
        return out

    def body_lines(container, parent) -> list[str]:
        out = []
        for child in container.iterchildren():
            if child.tag == qn("w:p"):
                out.append(Paragraph(child, parent).text)
            elif child.tag == qn("w:tbl"):
                out.extend(table_lines(Table(child, parent)))
        return out

    lines = body_lines(doc.element.body, doc)
    extra: list[str] = []
    seen = set()
    for section in doc.sections:
        for part in (section.header, section.footer):
            try:
                if part.is_linked_to_previous:
                    continue
                t = "\n".join(body_lines(part._element, part)).strip()
            except Exception:
                continue
            if t and t not in seen:
                seen.add(t)
                extra.append(t)
    text = "\n".join(lines).strip()
    if extra:
        text += "\n\n--- Kopf-/Fusszeilen ---\n" + "\n".join(extra)
    return {"text": text + "\n", "sections": []}


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def extract_text_file(data: bytes) -> dict:
    text = _decode_text(data).replace("\r\n", "\n").replace("\r", "\n")
    return {"text": text, "sections": []}


def extract(data: bytes, kind: str, max_pages: int = 200) -> dict:
    if kind == "pdf":
        result = extract_pdf(data, max_pages)
    elif kind == "xlsx":
        result = extract_xlsx(data)
    elif kind == "docx":
        result = extract_docx(data)
    elif kind in ("csv", "txt"):
        result = extract_text_file(data)
    else:
        raise ValueError(f"Nicht unterstuetzter Dateityp: {kind}")
    result.setdefault("extraction_method", {"xlsx": "openpyxl", "docx": "python-docx"}.get(kind, "text"))
    if len(result["text"]) > MAX_TEXT_CHARS:
        result["text"] = result["text"][:MAX_TEXT_CHARS]
        result["text_limit_reached"] = True
    return result


# ---------------------------------------------------------------------------
# Worker-Einstieg (eigener Prozess): Bytes auf stdin, JSON auf stdout
# ---------------------------------------------------------------------------

def _worker_main(argv: list[str]) -> int:
    kind = argv[1]
    max_pages = int(argv[2])
    max_mem_mb = int(argv[3]) if len(argv) > 3 else 0
    if max_mem_mb > 0:
        try:
            import resource
            limit = max_mem_mb * 1024 * 1024
            resource.setrlimit(resource.RLIMIT_AS, (limit, limit))
        except Exception:
            pass
    data = sys.stdin.buffer.read()
    try:
        out = {"ok": True, **extract(data, kind, max_pages)}
    except ExtractionLimit as e:
        out = {"ok": False, "error": str(e), "limit_exceeded": e.limit}
    except MemoryError:
        out = {"ok": False, "error": "Speicherlimit bei der Extraktion ueberschritten", "limit_exceeded": "memory"}
    except Exception as e:
        out = {"ok": False, "error": f"Extraktion fehlgeschlagen: {type(e).__name__}: {e}"}
    sys.stdout.write(json.dumps(out, ensure_ascii=False))
    sys.stdout.flush()
    return 0


if __name__ == "__main__":
    sys.exit(_worker_main(sys.argv))
