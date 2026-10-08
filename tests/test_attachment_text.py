import io
import os
import sys

import httpx
import pytest

sys.path.insert(0, os.path.dirname(__file__))

import zammad_mcp as zm  # noqa: E402
from zammad_mcp import attachment_text as at  # noqa: E402
from pdf_fixtures import simple_pdf, wps_like_pdf  # noqa: E402


def _xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "PO"
    ws.append(["Artikel", "Menge", "Preis"])
    ws.append(["Smart Lock", 10, 52.5])
    ws2 = wb.create_sheet("Notizen")
    ws2.append(["Lieferung KW 42"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx() -> bytes:
    from docx import Document
    doc = Document()
    doc.add_paragraph("Bestellung 4711")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Artikel", "Preis"
    t.cell(1, 0).text, t.cell(1, 1).text = "Smart Lock", "52.5"
    doc.add_paragraph("Ende")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Extraktion direkt
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("with_cmap,post_names,method", [
    (True, True, "pdfminer"),
    (False, True, "font_cmap"),
    (False, False, "gid_offset_heuristic"),
])
def test_pdf_methods(with_cmap, post_names, method):
    r = at.extract(wps_like_pdf(with_cmap, post_names), "pdf")
    assert r["extraction_method"] == method
    assert r["pages"] == 1
    assert r["text"].startswith("--- Seite 1 ---\n")
    # Tabellenzeile bleibt zusammen und in Lesereihenfolge
    assert "Smart Lock   10   52.5" in r["text"]
    assert r["text"].index("Delivery Note") < r["text"].index("Item") < r["text"].index("Cylinder")


def test_garbled_detection():
    assert at.is_garbled("(cid:39)(cid:72)(cid:79)")
    assert at.is_garbled("\x01\x02\x03ab")
    assert not at.is_garbled("Smart Lock 52.5 CHF")
    assert not at.is_garbled("")


def test_pdf_multi_page_sections():
    pdf = simple_pdf([[(50, 750, "Seite eins")], [(50, 750, "Preis 52.5")]])
    r = at.extract(pdf, "pdf")
    assert r["pages"] == 2
    assert "--- Seite 2 ---\nPreis 52.5" in r["text"]
    pos = r["text"].index("52.5")
    assert zm._section_at(r["sections"], pos)["page"] == 2


def test_pdf_page_limit():
    pdf = simple_pdf([[(50, 750, "x")]] * 3)
    with pytest.raises(at.ExtractionLimit) as e:
        at.extract(pdf, "pdf", max_pages=2)
    assert e.value.limit == "pages"


def test_xlsx_docx_csv_txt():
    x = at.extract(_xlsx(), "xlsx")
    assert "--- Blatt: PO ---" in x["text"] and "Smart Lock\t10\t52.5" in x["text"]
    assert zm._section_at(x["sections"], x["text"].index("KW 42"))["sheet"] == "Notizen"
    d = at.extract(_docx(), "docx")
    assert d["text"].splitlines()[:4] == ["Bestellung 4711", "Artikel\tPreis", "Smart Lock\t52.5", "Ende"]
    assert at.extract("a;b\r\nÄ;52,5\r\n".encode("cp1252"), "csv")["text"] == "a;b\nÄ;52,5\n"
    assert at.extract("Grüsse".encode("utf-8"), "txt")["text"] == "Grüsse"


def test_zip_bomb_limit(monkeypatch):
    monkeypatch.setattr(at, "MAX_ZIP_UNCOMPRESSED", 100)
    with pytest.raises(at.ExtractionLimit):
        at.extract(_xlsx(), "xlsx")


def test_detect_kind():
    assert at.detect_kind("a.PDF", "") == "pdf"
    assert at.detect_kind("scan", "application/pdf") == "pdf"
    assert at.detect_kind("bild.png", "image/png") is None
    assert at.detect_kind("a.doc", "application/msword") is None
    assert at.detect_kind("fake.pdf", "", b"not a pdf") is None
    assert at.detect_kind("fake.xlsx", "", b"%PDF-1.4 ...") == "pdf"


# ---------------------------------------------------------------------------
# Tools mit gemockter Zammad-API
# ---------------------------------------------------------------------------

def _att(i, name, data, ctype):
    return {"id": i, "filename": name, "size": str(len(data)), "preferences": {"Content-Type": ctype}}


@pytest.fixture
def zammad(monkeypatch):
    files = {
        101: ("po.pdf", wps_like_pdf(False, False), "application/pdf"),
        102: ("po.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        103: ("logo.png", b"\x89PNG....", "image/png"),
        104: ("notes.txt", b"nichts relevantes", "text/plain"),
        105: ("multi.pdf", simple_pdf([[(50, 750, "Seite eins")], [(50, 750, "Total 52.5 CHF")]]), "application/pdf"),
    }
    articles = {
        1: [{"id": 11, "ticket_id": 1, "attachments": [_att(i, *files[i]) for i in (101, 103)]},
            {"id": 12, "ticket_id": 1, "attachments": [_att(102, *files[102])]}],
        2: [{"id": 21, "ticket_id": 2, "attachments": [_att(i, *files[i]) for i in (104, 105, 101)]}],
    }
    tickets = {"1": {"id": 1, "number": "1001", "title": "QSD Smart Lock PO"},
               "2": {"id": 2, "number": "1002", "title": "QSD Smart Lock PO 2"}}

    def api_get(path, params=None):
        if path == "/tickets/search":
            return {"ticket_ids": [1, 2], "assets": {"Ticket": tickets}}
        if path.startswith("/ticket_articles/by_ticket/"):
            return articles[int(path.rsplit("/", 1)[1])]
        if path.startswith("/ticket_articles/"):
            aid = int(path.rsplit("/", 1)[1])
            return next(a for arts in articles.values() for a in arts if a["id"] == aid)
        raise AssertionError(path)

    downloads = []

    def download(ticket_id, article_id, attachment_id):
        downloads.append(attachment_id)
        name, data, ctype = files[attachment_id]
        return data, ctype

    def forbidden(*a, **k):
        raise AssertionError("schreibender Request")

    monkeypatch.setattr(zm, "api_get", api_get)
    monkeypatch.setattr(zm, "_download_attachment_limited", download)
    for m in ("post", "put", "patch", "delete"):
        monkeypatch.setattr(httpx, m, forbidden)
    zm._TEXT_CACHE.clear()
    return downloads


def test_get_attachment_text_paging(zammad):
    full = zm.get_attachment_text(1, 11, 101)
    assert full["extraction_method"] == "gid_offset_heuristic"
    assert full["pages"] == 1 and full["truncated"] is False and full["offset"] == 0
    assert full["returned_chars"] == full["total_chars"] == len(full["text"])
    p1 = zm.get_attachment_text(1, 11, 101, offset=0, max_chars=20)
    assert p1["truncated"] is True and p1["returned_chars"] == 20 and p1["next_offset"] == 20
    p2 = zm.get_attachment_text(1, 11, 101, offset=20, max_chars=10**6)
    assert p1["text"] + p2["text"] == full["text"] and p2["truncated"] is False
    assert zammad == [101]  # danach aus dem Cache
    assert zm.get_attachment_text(1, 11, 101, max_chars=10**6)["returned_chars"] <= zm.ATTACHMENT_TEXT_MAX_CHARS


def test_get_attachment_text_errors(zammad, monkeypatch):
    r = zm.get_attachment_text(1, 11, 103)
    assert "nicht unterstuetzt" in r["error"] and zammad == []
    r = zm.get_attachment_text(2, 11, 101)  # Artikel gehoert zu anderem Ticket
    assert "gehoert nicht zu Ticket 2" in r["error"] and zammad == []
    monkeypatch.setattr(zm, "ATTACHMENT_TEXT_MAX_BYTES", 10)
    r = zm.get_attachment_text(1, 12, 102)
    assert r["limit_exceeded"] == "size" and zammad == []


def test_page_limit_and_timeout(zammad, monkeypatch):
    monkeypatch.setattr(zm, "ATTACHMENT_TEXT_MAX_PDF_PAGES", 1)
    r = zm.get_attachment_text(2, 21, 105)
    assert r["limit_exceeded"] == "pages"
    monkeypatch.setattr(zm, "ATTACHMENT_TEXT_TIMEOUT", 0.001)
    r = zm.get_attachment_text(1, 12, 102)
    assert r["limit_exceeded"] == "timeout"


def test_find_in_attachments(zammad):
    r = zm.find_in_attachments("QSD Smart Lock PO", "52.5", context_chars=20)
    assert r["tickets_searched"] == 2
    found = {(m["attachment_id"], m.get("page"), m.get("sheet")) for m in r["matches"]}
    assert found == {(101, 1, None), (102, None, "PO"), (105, 2, None)}
    # 101 haengt an zwei Tickets, wird aber nur einmal verarbeitet
    assert sorted(zammad) == [101, 102, 104, 105]
    hit = next(m for m in r["matches"] if m["attachment_id"] == 101)
    assert hit["ticket_number"] == "1001" and hit["article_id"] == 11 and "52.5" in hit["context"]
    assert hit["extraction_method"] == "gid_offset_heuristic"
    assert [a["filename"] for a in r["attachments_without_match"]] == ["notes.txt"]
    assert r["skipped_unsupported_count"] == 1 and r["skipped_unsupported"][0]["filename"] == "logo.png"
    assert r["attachments_checked"] == 4


def test_find_substring_is_literal_and_case_insensitive(zammad):
    r = zm.find_in_attachments("x", "SMART LOCK")
    assert {m["attachment_id"] for m in r["matches"]} == {101, 102}
    r = zm.find_in_attachments("x", "52.5.")  # Punkt ist kein Wildcard
    assert r["match_count"] == 0


def test_find_regex(zammad):
    r = zm.find_in_attachments("x", r"\b52[.,]5\b", regex=True)
    assert len(r["matches"]) == 3
    with pytest.raises(ValueError):
        zm.find_in_attachments("x", "a" * 201, regex=True)
    with pytest.raises(ValueError):
        zm.find_in_attachments("x", "(", regex=True)


def test_find_regex_timeout(zammad, monkeypatch):
    monkeypatch.setattr(zm, "REGEX_TIMEOUT", 0.2)
    monkeypatch.setattr(zm, "_extract_attachment", lambda *a, **k: {
        "ok": True, "text": "x" * 5000, "sections": [], "extraction_method": "text"})
    r = zm.find_in_attachments("x", r"(x+x+)+y", regex=True)
    assert r["errors"] and all(e["limit_exceeded"] == "regex_timeout" for e in r["errors"])


def test_tools_registered():
    import asyncio
    names = {t.name for t in asyncio.run(zm.list_tools())}
    assert {"get_attachment_text", "find_in_attachments"} <= names
    assert {"get_attachment_text", "find_in_attachments"} <= set(zm.TOOL_FUNCS)


class _FakeStream:
    def __init__(self, chunks, headers):
        self.chunks, self.headers = chunks, headers

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def raise_for_status(self):
        pass

    def iter_bytes(self):
        yield from self.chunks


def test_download_size_limit(monkeypatch):
    monkeypatch.setattr(zm, "ZAMMAD_URL", "https://zammad.invalid")
    monkeypatch.setattr(zm, "ATTACHMENT_TEXT_MAX_BYTES", 10)
    calls = []

    def stream(method, url, **kw):
        calls.append((method, url))
        return _FakeStream([b"12345", b"67890", b"x"], {})

    monkeypatch.setattr(httpx, "stream", stream)
    with pytest.raises(zm._AttachmentLimit):
        zm._download_attachment_limited(1, 2, 3)
    assert calls == [("GET", "https://zammad.invalid/api/v1/ticket_attachment/1/2/3")]
    monkeypatch.setattr(httpx, "stream", lambda *a, **k: _FakeStream([b"ok"], {"content-length": "999"}))
    with pytest.raises(zm._AttachmentLimit):
        zm._download_attachment_limited(1, 2, 3)
    monkeypatch.setattr(httpx, "stream", lambda *a, **k: _FakeStream([b"ok"], {"content-type": "text/plain"}))
    assert zm._download_attachment_limited(1, 2, 3) == (b"ok", "text/plain")


def test_call_tool_json(zammad):
    import asyncio
    import json
    out = asyncio.run(zm.call_tool("get_attachment_text", {"ticket_id": 1, "article_id": 12, "attachment_id": 102}))
    data = json.loads(out[0].text)
    assert data["extraction_method"] == "openpyxl" and "Smart Lock\t10\t52.5" in data["text"]
